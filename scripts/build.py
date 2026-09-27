# -*- coding: utf-8 -*-
"""
data/*.csv からデータベースと配布用ファイルを再生成する。

    python3 scripts/build.py

入力  : data/vendors.csv, categories.csv, vendor_categories.csv,
        attributes.csv, vendor_attributes.csv, exhibitions.csv, excluded.csv
出力  : db/ipai.sqlite            SQLite（SQLで引く用）
        export/vendors_full.csv   1行1社のフラットCSV（Excel用）
        export/vendors.json       Webツール・記事生成用
        export/data.js            Webツール（Artifact）差し替え用
        export/exhibitors_2026.csv 展示会出展社のみ
        export/vendors.xlsx       Excel（openpyxlがある場合のみ）

CSVを直接編集したあとにこれを流せば、すべての形式が揃う。
"""

import csv
import json
import pathlib
import sqlite3
import sys

ROOT = pathlib.Path(__file__).resolve().parents[1]
DATA, DB, EXPORT = ROOT / "data", ROOT / "db", ROOT / "export"
DBFILE = DB / "ipai.sqlite"

SCHEMA = """
DROP VIEW  IF EXISTS v_vendor_full;
DROP TABLE IF EXISTS vendor_attributes;
DROP TABLE IF EXISTS attributes;
DROP TABLE IF EXISTS vendor_categories;
DROP TABLE IF EXISTS exhibitions;
DROP TABLE IF EXISTS excluded;
DROP TABLE IF EXISTS vendors;
DROP TABLE IF EXISTS categories;

CREATE TABLE categories(
  cat_id TEXT PRIMARY KEY, phase TEXT, title TEXT, short TEXT, subtitle TEXT, sort INTEGER);

CREATE TABLE attributes(
  attr_id TEXT PRIMARY KEY, label TEXT, short TEXT, definition TEXT, sort INTEGER);

CREATE TABLE vendors(
  vendor_id TEXT PRIMARY KEY, name TEXT NOT NULL UNIQUE, region TEXT, country TEXT,
  url TEXT, description TEXT, deployment TEXT,
  is_mcp INTEGER DEFAULT 0, is_new_2026 INTEGER DEFAULT 0,
  status TEXT DEFAULT 'active', first_listed TEXT, updated_at TEXT, note TEXT);

CREATE TABLE vendor_categories(
  vendor_id TEXT, cat_id TEXT,
  role TEXT DEFAULT '主',          -- 主=カオスマップ上の代表機能 / 副=確認できた対応機能
  source TEXT, checked_at TEXT,    -- 根拠と確認日
  PRIMARY KEY(vendor_id, cat_id),
  FOREIGN KEY(vendor_id) REFERENCES vendors(vendor_id),
  FOREIGN KEY(cat_id)   REFERENCES categories(cat_id));

CREATE TABLE vendor_attributes(
  vendor_id TEXT, attr_id TEXT,
  value TEXT,                      -- あり / なし / 未確認（行がない組み合わせも未確認）
  source TEXT, checked_at TEXT, note TEXT,
  PRIMARY KEY(vendor_id, attr_id),
  FOREIGN KEY(vendor_id) REFERENCES vendors(vendor_id),
  FOREIGN KEY(attr_id)   REFERENCES attributes(attr_id));

CREATE TABLE exhibitions(
  exh_id TEXT PRIMARY KEY, event TEXT, year INTEGER, venue TEXT, period TEXT,
  vendor_id TEXT, vendor_name TEXT, booth TEXT, exhibit_name TEXT, exhibit_note TEXT,
  FOREIGN KEY(vendor_id) REFERENCES vendors(vendor_id));

CREATE TABLE excluded(
  excluded_id TEXT PRIMARY KEY, name TEXT, kind TEXT, reason TEXT, checked_at TEXT);

CREATE INDEX idx_vc_cat    ON vendor_categories(cat_id);
CREATE INDEX idx_ex_vendor ON exhibitions(vendor_id);
CREATE INDEX idx_v_region  ON vendors(region);

CREATE VIEW v_vendor_full AS
SELECT v.vendor_id, v.name, v.region, v.country, v.deployment,
       v.is_mcp, v.is_new_2026, v.url, v.description,
       (SELECT group_concat(c.title, ' / ')
          FROM vendor_categories vc JOIN categories c ON c.cat_id = vc.cat_id
         WHERE vc.vendor_id = v.vendor_id ORDER BY c.sort) AS categories,
       (SELECT group_concat(vc.cat_id, ',')
          FROM vendor_categories vc WHERE vc.vendor_id = v.vendor_id) AS cat_ids,
       (SELECT count(*) FROM vendor_categories vc WHERE vc.vendor_id = v.vendor_id) AS cat_count,
       (SELECT group_concat(c.title, ' / ')
          FROM vendor_categories vc JOIN categories c ON c.cat_id = vc.cat_id
         WHERE vc.vendor_id = v.vendor_id AND vc.role = '主' ORDER BY c.sort) AS primary_categories,
       (SELECT group_concat(vc.cat_id, ',')
          FROM vendor_categories vc WHERE vc.vendor_id = v.vendor_id AND vc.role = '主') AS primary_cat_ids,
       COALESCE((SELECT value FROM vendor_attributes a WHERE a.vendor_id = v.vendor_id AND a.attr_id = 'ja'), '未確認') AS ja,
       COALESCE((SELECT value FROM vendor_attributes a WHERE a.vendor_id = v.vendor_id AND a.attr_id = 'onprem'), '未確認') AS onprem,
       COALESCE((SELECT value FROM vendor_attributes a WHERE a.vendor_id = v.vendor_id AND a.attr_id = 'no_training'), '未確認') AS no_training,
       COALESCE((SELECT value FROM vendor_attributes a WHERE a.vendor_id = v.vendor_id AND a.attr_id = 'trial'), '未確認') AS trial,
       COALESCE((SELECT value FROM vendor_attributes a WHERE a.vendor_id = v.vendor_id AND a.attr_id = 'pricing'), '未確認') AS pricing,
       e.booth, e.exhibit_name, e.exhibit_note, e.event AS exhibited_at
  FROM vendors v
  LEFT JOIN exhibitions e ON e.vendor_id = v.vendor_id
 WHERE v.status = 'active';
"""

TABLES = {
    "categories": ["cat_id", "phase", "title", "short", "subtitle", "sort"],
    "vendors": ["vendor_id", "name", "region", "country", "url", "description",
                "deployment", "is_mcp", "is_new_2026", "status", "first_listed",
                "updated_at", "note"],
    "vendor_categories": ["vendor_id", "cat_id", "role", "source", "checked_at"],
    "attributes": ["attr_id", "label", "short", "definition", "sort"],
    "vendor_attributes": ["vendor_id", "attr_id", "value", "source", "checked_at", "note"],
    "exhibitions": ["exh_id", "event", "year", "venue", "period", "vendor_id",
                    "vendor_name", "booth", "exhibit_name", "exhibit_note"],
    "excluded": ["excluded_id", "name", "kind", "reason", "checked_at"],
}


def read(name):
    p = DATA / f"{name}.csv"
    if not p.exists():
        sys.exit(f"[エラー] {p} がありません")
    with p.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def validate(tbl):
    """壊れたデータのまま出力しないための検査"""
    errs, warns = [], []
    ids = {v["vendor_id"] for v in tbl["vendors"]}
    cids = {c["cat_id"] for c in tbl["categories"]}
    names = [v["name"] for v in tbl["vendors"]]

    if len(ids) != len(tbl["vendors"]):
        errs.append("vendors.csv の vendor_id に重複がある")
    if len(set(names)) != len(names):
        dup = {n for n in names if names.count(n) > 1}
        errs.append(f"vendors.csv の name に重複がある: {sorted(dup)}")
    for r in tbl["vendor_categories"]:
        if r["vendor_id"] not in ids:
            errs.append(f"vendor_categories: 存在しない vendor_id {r['vendor_id']}")
        if r["cat_id"] not in cids:
            errs.append(f"vendor_categories: 存在しない cat_id {r['cat_id']}")
    for r in tbl["vendor_categories"]:
        if r.get("role", "主") not in ("主", "副"):
            errs.append(f"vendor_categories: role は 主/副 のみ（{r['vendor_id']} {r['cat_id']}: {r.get('role')}）")
        if r.get("role") == "副" and not r.get("source"):
            warns.append(f"根拠なしの副工程: {r['vendor_id']} {r['cat_id']}")
    for r in tbl["exhibitions"]:
        if r["vendor_id"] and r["vendor_id"] not in ids:
            errs.append(f"exhibitions: 存在しない vendor_id {r['vendor_id']}")

    # 導入条件：値は3種類に限り、「あり」「なし」には根拠を必須にする
    aids = {a["attr_id"] for a in tbl["attributes"]}
    seen = set()
    for r in tbl["vendor_attributes"]:
        key = (r["vendor_id"], r["attr_id"])
        if r["vendor_id"] not in ids:
            errs.append(f"vendor_attributes: 存在しない vendor_id {r['vendor_id']}")
        if r["attr_id"] not in aids:
            errs.append(f"vendor_attributes: 存在しない attr_id {r['attr_id']}")
        if key in seen:
            errs.append(f"vendor_attributes: 重複 {r['vendor_id']} {r['attr_id']}")
        seen.add(key)
        if r["value"] not in ("あり", "なし", "未確認"):
            errs.append(f"vendor_attributes: value は あり/なし/未確認 のみ（{r['vendor_id']} {r['attr_id']}: {r['value']}）")
        if r["value"] in ("あり", "なし") and not r.get("source"):
            errs.append(f"vendor_attributes: 根拠のない判定 {r['vendor_id']} {r['attr_id']}")

    linked = {r["vendor_id"] for r in tbl["vendor_categories"]}
    for v in tbl["vendors"]:
        if v["vendor_id"] not in linked and v.get("status", "active") == "active":
            warns.append(f"カテゴリ未設定: {v['name']}")
        if not v["url"]:
            warns.append(f"URL未確認: {v['name']}")
        if not v["description"]:
            warns.append(f"紹介文なし: {v['name']}")
    return errs, warns


def build_sqlite(tbl):
    DB.mkdir(exist_ok=True)
    con = sqlite3.connect(DBFILE)
    con.executescript(SCHEMA)
    for t, cols in TABLES.items():
        con.executemany(
            f"INSERT INTO {t}({','.join(cols)}) VALUES({','.join('?' * len(cols))})",
            [[r.get(c, "") for c in cols] for r in tbl[t]])
    con.commit()
    return con


def export_csv(con):
    EXPORT.mkdir(exist_ok=True)
    cur = con.execute("SELECT * FROM v_vendor_full ORDER BY name")
    cols = [d[0] for d in cur.description]
    rows = cur.fetchall()
    with (EXPORT / "vendors_full.csv").open("w", encoding="utf-8-sig", newline="") as f:
        cw = csv.writer(f)
        cw.writerow(cols)
        cw.writerows(rows)

    cur = con.execute("""
        SELECT e.booth, v.name, e.exhibit_name, v.region,
               (SELECT group_concat(c.title, ' / ') FROM vendor_categories vc
                  JOIN categories c ON c.cat_id = vc.cat_id
                 WHERE vc.vendor_id = v.vendor_id ORDER BY c.sort) AS categories,
               v.url, e.exhibit_note
          FROM exhibitions e JOIN vendors v ON v.vendor_id = e.vendor_id
         ORDER BY e.booth""")
    with (EXPORT / "exhibitors_2026.csv").open("w", encoding="utf-8-sig", newline="") as f:
        cw = csv.writer(f)
        cw.writerow([d[0] for d in cur.description])
        cw.writerows(cur.fetchall())
    return len(rows)


def export_json(con, tbl):
    cats = [dict(cat_id=r["cat_id"], phase=r["phase"], title=r["title"], short=r["short"],
                 subtitle=r["subtitle"], sort=int(r["sort"])) for r in tbl["categories"]]
    attrs = [dict(attr_id=r["attr_id"], label=r["label"], short=r["short"],
                  definition=r["definition"], sort=int(r["sort"])) for r in tbl["attributes"]]
    attr_ids = [a["attr_id"] for a in attrs]
    con.row_factory = sqlite3.Row
    vendors = []
    for r in con.execute("SELECT * FROM v_vendor_full ORDER BY cat_ids, name"):
        d = {
            "vendor_id": r["vendor_id"], "name": r["name"], "region": r["region"],
            "country": r["country"], "url": r["url"], "desc": r["description"],
            "deployment": r["deployment"], "cats": sorted((r["cat_ids"] or "").split(",")),
            "primary": sorted((r["primary_cat_ids"] or "").split(",")) if r["primary_cat_ids"] else [],
            "mcp": bool(r["is_mcp"]), "new": bool(r["is_new_2026"]),
            # 導入条件は「未確認」を省き、あり／なしだけを持たせる
            "attrs": {a: r[a] for a in attr_ids if r[a] != "未確認"},
        }
        if r["booth"]:
            d["booth"] = r["booth"]
            d["exhibit"] = r["exhibit_name"]
            d["note"] = r["exhibit_note"]
        elif r["exhibited_at"]:
            d["booth"] = "—"
            d["exhibit"] = r["exhibit_name"]
            d["note"] = r["exhibit_note"]
        vendors.append(d)
    con.row_factory = None

    updated = max((v["updated_at"] for v in tbl["vendors"] if v["updated_at"]), default="")
    data = {
        "meta": {
            "updated": updated, "vendors": len(vendors), "cats": len(cats),
            "exhibitors": sum(1 for v in vendors if "booth" in v),
            "excluded": len(tbl["excluded"]),
            "fair": "2026 知財・情報フェア＆コンファレンス（2026年9月16〜18日・東京ビッグサイト東3ホール）",
            "basis": "AI・生成AIを組み込んだ製品またはサービスを公開情報で確認できたベンダーに限る",
        },
        "categories": cats,
        "attributes": attrs,
        "vendors": vendors,
        "excluded": [dict(name=r["name"], kind=r["kind"], reason=r["reason"])
                     for r in tbl["excluded"]],
    }
    body = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    (EXPORT / "vendors.json").write_text(body, encoding="utf-8")
    # Webツール（Artifact）にそのまま差し替えられる形
    (EXPORT / "data.js").write_text("const DB=" + body + ";", encoding="utf-8")
    return data


def export_xlsx(con):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None
    cur = con.execute("SELECT * FROM v_vendor_full ORDER BY name")
    cols = [d[0] for d in cur.description]
    rows = cur.fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "vendors"
    ws.append(cols)
    for r in rows:
        ws.append(list(r))
    head = PatternFill("solid", fgColor="1F3B57")
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
    widths = {"name": 26, "region": 11, "country": 13, "deployment": 22,
              "url": 34, "description": 78, "categories": 40, "cat_ids": 14,
              "booth": 9, "exhibit_name": 28, "exhibit_note": 60}
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 12)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows)+1}"

    # 導入条件の根拠（1行1判定）
    cur = con.execute("""
        SELECT v.name, a.label, va.value, va.note, va.source, va.checked_at
          FROM vendor_attributes va
          JOIN vendors v ON v.vendor_id = va.vendor_id
          JOIN attributes a ON a.attr_id = va.attr_id
         WHERE v.status = 'active'
         ORDER BY v.name, a.sort""")
    ev_cols = ["name", "attribute", "value", "note", "source", "checked_at"]
    ws2 = wb.create_sheet("attributes_evidence")
    ws2.append(ev_cols)
    ev_rows = cur.fetchall()
    for r in ev_rows:
        ws2.append(list(r))
    for c in range(1, len(ev_cols) + 1):
        cell = ws2.cell(1, c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head
    for i, w_ in enumerate([26, 30, 9, 60, 50, 12], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w_
    ws2.freeze_panes = "B2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(ev_cols))}{len(ev_rows)+1}"
    wb.save(EXPORT / "vendors.xlsx")
    return len(rows)


def main():
    tbl = {t: read(t) for t in TABLES}
    errs, warns = validate(tbl)
    if errs:
        print("■ エラー（出力を中止）")
        for e in errs:
            print("  -", e)
        sys.exit(1)

    con = build_sqlite(tbl)
    n = export_csv(con)
    data = export_json(con, tbl)
    xlsx = export_xlsx(con)
    con.close()

    m = data["meta"]
    print("■ 再生成しました")
    print(f"  ベンダー {m['vendors']}社 / カテゴリ {m['cats']} / "
          f"フェア出展 {m['exhibitors']}社 / 対象外 {m['excluded']}件")
    counts = "・".join(
        f"{a['short']} {sum(1 for v in data['vendors'] if v['attrs'].get(a['attr_id']) == 'あり')}"
        for a in data["attributes"])
    print(f"  導入条件（あり）: {counts}")
    print(f"  db/ipai.sqlite")
    print(f"  export/vendors_full.csv（{n}行）")
    print(f"  export/exhibitors_2026.csv")
    print(f"  export/vendors.json")
    print(f"  export/data.js（Webツール差し替え用）")
    print("  export/vendors.xlsx" if xlsx else
          "  （xlsxは openpyxl 未導入のためスキップ: pip install openpyxl）")

    if warns:
        print(f"\n■ 確認したい箇所 {len(warns)}件")
        for w_ in warns[:12]:
            print("  -", w_)
        if len(warns) > 12:
            print(f"  ... ほか{len(warns)-12}件")


if __name__ == "__main__":
    main()
